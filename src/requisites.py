from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any



@dataclass
class Requisites:
    company_name: str = ""
    inn: str = ""
    kpp: str = ""
    ogrn: str = ""
    legal_address: str = ""
    bank_name: str = ""
    bik: str = ""
    checking_account: str = ""
    correspondent_account: str = ""
    director_name: str = ""
    acting_on: str = "Устава"

    def to_context(self) -> dict[str, str]:
        contract_date = datetime.now().strftime("%d.%m.%Y")
        preamble_client = self._build_preamble(contract_date)
        signature_block_client = self._build_signature_block()

        context = {
            "contract_date": contract_date,
            "company_name": self.company_name,
            "inn": self.inn,
            "kpp": self.kpp,
            "ogrn": self.ogrn,
            "legal_address": self.legal_address,
            "bank_name": self.bank_name,
            "bik": self.bik,
            "checking_account": self.checking_account,
            "correspondent_account": self.correspondent_account,
            "director_name": self.director_name,
            "acting_on": self.acting_on,
            "preamble_client": preamble_client,
            "signature_block_client": signature_block_client,
        }

        return {k: (v or "") for k, v in context.items()}

    def _build_preamble(self, contract_date: str) -> str:
        company = self.company_name or "________________"
        director = self.director_name or "________________"
        acting_on = self.acting_on or "________________"

        return (
            f"{company}, именуемое в дальнейшем «Заказчик», в лице {director}, "
            f"действующего на основании {acting_on}, с одной стороны, заключило настоящий "
            f"договор от {contract_date}."
        )

    def _build_signature_block(self) -> str:
        company = self.company_name or "________________"
        director = self.director_name or "________________"

        lines = [
            f"{company}",
            f"ИНН/КПП: {self.inn}/{self.kpp}" if self.inn or self.kpp else "ИНН/КПП: ",
            f"ОГРН: {self.ogrn}" if self.ogrn else "ОГРН: ",
            f"Юр. адрес: {self.legal_address}" if self.legal_address else "Юр. адрес: ",
            f"Банк: {self.bank_name}" if self.bank_name else "Банк: ",
            f"БИК: {self.bik}" if self.bik else "БИК: ",
            (
                f"р/с: {self.checking_account}, к/с: {self.correspondent_account}"
                if self.checking_account or self.correspondent_account
                else "р/с: , к/с: "
            ),
            f"Подпись: ____________ /{director}/",
        ]
        return "\n".join(lines)


def parse_uploaded_requisites(file_name: str, file_bytes: bytes) -> Requisites:
    suffix = file_name.lower().rsplit(".", maxsplit=1)[-1] if "." in file_name else ""

    if suffix == "json":
        data = json.loads(file_bytes.decode("utf-8"))
        return parse_from_json_dict(data)
    if suffix in {"txt", "csv"}:
        text = file_bytes.decode("utf-8", errors="ignore")
        return parse_from_text(text)
    if suffix == "docx":
        text = extract_text_from_docx_bytes(file_bytes)
        return parse_from_text(text)

    text = file_bytes.decode("utf-8", errors="ignore")
    return parse_from_text(text)


def parse_requisites_from_url(url: str, timeout_s: int = 20) -> Requisites:
    import requests
    from bs4 import BeautifulSoup

    response = requests.get(url, timeout=timeout_s)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "lxml")
    text = soup.get_text("\n", strip=True)
    return parse_from_text(text)


def parse_from_json_dict(data: dict[str, Any]) -> Requisites:
    norm = {str(k).lower().strip(): ("" if v is None else str(v).strip()) for k, v in data.items()}

    def pick(*keys: str) -> str:
        for key in keys:
            if key in norm and norm[key]:
                return norm[key]
        return ""

    return Requisites(
        company_name=pick("company_name", "company", "name", "название", "наименование", "название юл"),
        inn=pick("inn", "инн"),
        kpp=pick("kpp", "кпп"),
        ogrn=pick("ogrn", "огрн"),
        legal_address=pick("legal_address", "address", "юридический адрес", "юр адрес", "адрес"),
        bank_name=pick("bank_name", "bank", "банк"),
        bik=pick("bik", "бик"),
        checking_account=pick("checking_account", "account", "р/с", "расчетный счет"),
        correspondent_account=pick("correspondent_account", "к/с", "корреспондентский счет"),
        director_name=pick("director_name", "director", "генеральный директор", "директор", "подписант"),
        acting_on=pick("acting_on", "основание", "действует на основании") or "Устава",
    )


def parse_from_text(text: str) -> Requisites:
    cleaned = re.sub(r"\u00a0", " ", text)

    company_name = _find_first(
        cleaned,
        [
            r"(?:ООО|АО|ПАО|ЗАО|ИП)\s+[«\"“]?[A-ЯA-Za-z0-9\s\-_.]+[»\"”]?",
            r"Наименование\s*[:\-]\s*(.+)",
        ],
    )

    return Requisites(
        company_name=_cleanup_company_name(company_name),
        inn=_digits_only(_find_first(cleaned, [r"ИНН\s*[:\-]?\s*([0-9]{10,12})", r"\b([0-9]{10,12})\b"])),
        kpp=_digits_only(_find_first(cleaned, [r"КПП\s*[:\-]?\s*([0-9]{9})"])),
        ogrn=_digits_only(_find_first(cleaned, [r"ОГРН\s*[:\-]?\s*([0-9]{13,15})"])),
        legal_address=_normalize_space(
            _find_first(
                cleaned,
                [
                    r"Юрид(?:ический|ическ\.)\s*адрес\s*[:\-]\s*(.+)",
                    r"Юр\.?\s*адрес\s*[:\-]\s*(.+)",
                ],
            )
        ),
        bank_name=_normalize_space(_find_first(cleaned, [r"Банк\s*[:\-]\s*(.+)"])),
        bik=_digits_only(_find_first(cleaned, [r"БИК\s*[:\-]?\s*([0-9]{9})"])),
        checking_account=_digits_only(
            _find_first(cleaned, [r"(?:р/с|расч(?:етный)?\s*счет)\s*[:\-]?\s*([0-9]{20})"])
        ),
        correspondent_account=_digits_only(
            _find_first(cleaned, [r"(?:к/с|корр(?:еспондентский)?\s*счет)\s*[:\-]?\s*([0-9]{20})"])
        ),
        director_name=_normalize_space(
            _find_first(
                cleaned,
                [
                    r"Генеральный\s*директор\s*[:\-]\s*([A-ЯЁA-Z][^\n,]+)",
                    r"Директор\s*[:\-]\s*([A-ЯЁA-Z][^\n,]+)",
                    r"В\s*лице\s*([A-ЯЁA-Z][^\n,]+)",
                ],
            )
        ),
        acting_on=_normalize_space(
            _find_first(cleaned, [r"(?:действующ(?:ий|его)\s*на\s*основании|Основание)\s*[:\-]?\s*([^\n.]+)"])
        )
        or "Устава",
    )


def extract_text_from_docx_bytes(file_bytes: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(file_bytes))
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt:
                    parts.append(txt)

    return "\n".join(parts)


def _find_first(text: str, patterns: list[str]) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if not match:
            continue
        if match.groups():
            return match.group(1).strip()
        return match.group(0).strip()
    return ""


def _digits_only(value: str) -> str:
    return re.sub(r"\D", "", value or "")


def _normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def _cleanup_company_name(value: str) -> str:
    value = _normalize_space(value)
    return value.strip("-:;,. ")
